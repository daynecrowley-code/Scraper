"""Accumulates rows the agent records and writes them to xlsx/csv."""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

META_COLS = ["confidence", "evidence", "source_url"]


class ResultSet:
    def __init__(self) -> None:
        self.columns: list[str] = []
        self.rows: list[dict] = []
        self.notes: str = ""

    def define_columns(self, columns: list[str]) -> str:
        if self.rows:
            return "ERROR: columns can't change after rows were recorded"
        self.columns = [c for c in dict.fromkeys(columns) if c not in META_COLS]
        return f"Columns set: {self.columns + META_COLS}"

    def add_rows(self, rows: list[dict]) -> str:
        if not self.columns:
            return "ERROR: call define_columns first"
        added = dups = 0
        seen = {self._key(r) for r in self.rows}
        for r in rows:
            clean = {c: r.get(c, "") for c in self.columns + META_COLS}
            if self._key(clean) in seen:
                dups += 1
                continue
            seen.add(self._key(clean))
            self.rows.append(clean)
            added += 1
        return f"Recorded {added} rows ({dups} duplicates skipped). Total: {len(self.rows)}"

    def _key(self, r: dict) -> tuple:
        return tuple(str(r.get(c, "")).strip().lower() for c in self.columns)

    def save(self, path: Path) -> Path:
        path.parent.mkdir(parents=True, exist_ok=True)
        headers = self.columns + META_COLS
        if path.suffix == ".csv":
            with path.open("w", newline="", encoding="utf-8") as f:
                w = csv.DictWriter(f, fieldnames=headers)
                w.writeheader()
                w.writerows(self.rows)
            return path
        wb = Workbook()
        ws = wb.active
        ws.title = "Results"
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in self.rows:
            ws.append([r[h] for h in headers])
        ws.freeze_panes = "A2"
        for i, h in enumerate(headers, 1):
            width = max([len(h)] + [len(str(r[h])) for r in self.rows]) if self.rows else len(h)
            ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 60)
        notes = wb.create_sheet("Notes")
        notes.append(["Agent summary / caveats"])
        notes["A1"].font = Font(bold=True)
        for line in (self.notes or "").splitlines():
            notes.append([line])
        wb.save(path)
        return path
