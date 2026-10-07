from openpyxl import load_workbook

from scraper_agent.fetch import html_to_text
from scraper_agent.results import ResultSet


def test_html_to_text_tables_and_links():
    t = html_to_text("<script>x</script><table><tr><th>Item</th><th>Price</th></tr>"
                     "<tr><td>Cumin</td><td>$4/lb</td></tr></table><a href='/a'>More</a>")
    assert "Item | Price" in t and "Cumin | $4/lb" in t and "More </a>" in t
    assert "x" not in t.split()


def test_results_dedupe_and_xlsx(tmp_path):
    rs = ResultSet()
    assert rs.add_rows([{}]).startswith("ERROR")
    rs.define_columns(["event", "city"])
    msg = rs.add_rows([{"event": "NPE", "city": "Orlando", "confidence": "likely"},
                       {"event": "npe", "city": "orlando"}])
    assert "1 rows" in msg and "1 duplicates" in msg
    rs.notes = "ok"
    p = rs.save(tmp_path / "o.xlsx")
    ws = load_workbook(p)["Results"]
    assert [c.value for c in ws[1]] == ["event", "city", "confidence", "evidence", "source_url"]
    assert ws.max_row == 2
